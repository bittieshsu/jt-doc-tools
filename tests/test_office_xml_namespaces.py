"""文件翻譯寫回檔案時，**命名空間的前綴與宣告要照原檔**。

## 由來（v1.16.9，客戶回報「xlsx 翻譯預覽可以，下載後變空白」）

ElementTree 輸出時會把沒註冊的前綴改名（`x14ac` → `ns3`）、丟掉沒用到的宣告
（`xr2` / `xr3`）。XML 本身仍然合法，但 Office 有兩個地方是**用前綴的名字**指的：

* `mc:Ignorable="x14ac xr xr2 xr3"`（Excel 存的工作表、Word 存的主文件都有）
* `<mc:Choice Requires="wps">`（Word 的文字方塊）

名字指向不存在的宣告 → Excel 判定內容有問題、修復時把資料清掉 → **空白試算表**。
LibreOffice 不理這些屬性，所以用它畫的預覽完全正常。

**這個洞從文件翻譯上線就在**：我們手上的試算表樣本都是 OxOffice / LibreOffice
存的，沒有 `mc:Ignorable`，所有既有測試都綠。所以這裡的素材刻意做成
**Excel / Word 存出來的樣子**。

## 判準

* 被 `Ignorable` / `Requires` 點名的前綴**每一個都有宣告**
* 每個網址用的前綴**跟原檔一樣**（不可以出現 `ns0` 這種改名）
* `standalone="yes"` 留著
* 譯文真的寫進去了、openpyxl 讀得回來
* 全行程共用的 `ET._namespace_map` 用完要還原（它影響同一個行程裡所有的 XML 輸出）
"""
from __future__ import annotations

import io
import re
import threading
import zipfile
from xml.etree import ElementTree as ET

import pytest
from defusedxml.ElementTree import fromstring

from app.core import office_text_map as M

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_EXCEL_ROOT = (
    '<worksheet xmlns="' + _MAIN + '" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'mc:Ignorable="x14ac xr xr2 xr3" '
    'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" '
    'xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" '
    'xmlns:xr2="http://schemas.microsoft.com/office/spreadsheetml/2015/revision2" '
    'xmlns:xr3="http://schemas.microsoft.com/office/spreadsheetml/2016/revision3" '
    'xr:uid="{00000000-0001-0000-0000-000000000000}">'
)


def _excel_like_xlsx() -> bytes:
    """openpyxl 做一份，再把工作表的根元素換成 Excel 存檔時的樣子。

    `xr2` / `xr3` 在內文**一次都沒用到** —— 那正是 ElementTree 會丟掉的宣告。
    """
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Quarterly Report"
    ws["A2"] = "Security requirements for vendors"
    ws["B2"] = "Owner"
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<worksheet\b[^>]*>", _EXCEL_ROOT, text, count=1)
                text = re.sub(r"<row\b", '<row x14ac:dyDescent="0.25"', text)
                # openpyxl 不寫 XML 宣告；Excel 寫的是這一行
                text = re.sub(r"^<\?xml[^>]*\?>\s*", "", text)
                text = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
                        + text)
                assert "x14ac:dyDescent" in text, "素材沒做成：x14ac 沒有被用到"
                data = text.encode("utf-8")
            z.writestr(info, data)
    return out.getvalue()


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _word_like_docx() -> bytes:
    """Word 存出來的主文件：`mc:Ignorable` ＋ 文字方塊的 `Requires="wps"`。"""
    tb = ('<w:txbxContent><w:p><w:r><w:t>Text in a box</w:t></w:r></w:p>'
          '</w:txbxContent>')
    doc = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        '<w:document '
        'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        f'xmlns:w="{_W}" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
        'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'mc:Ignorable="w14 w15 wp14">'
        '<w:body>'
        '<w:p w14:paraId="1A2B3C4D" w14:textId="77777777">'
        '<w:r><w:t>Quarterly Report</w:t></w:r></w:p>'
        '<w:p w14:paraId="1A2B3C4E" w14:textId="77777777"><w:r>'
        '<mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor><wps:wsp><wps:txbx>'
        + tb +
        '</wps:txbx></wps:wsp></wp:anchor></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><v:shape o:allowincell="f"><v:textbox>'
        + tb +
        '</v:textbox></v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r></w:p>'
        '</w:body></w:document>'
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", doc)
    return buf.getvalue()


def _translate(data: bytes, ext: str) -> tuple[bytes, dict]:
    units, state = M.extract_units(data, ext)
    assert units, "素材裡沒有抽到任何段落"
    out = M.rebuild(state, {i: f"譯：{u.text}" for i, u in enumerate(units)}, units)
    return out, state


def _decls(xml: bytes) -> dict[str, str]:
    """前綴 → 網址（整份文件的宣告）。"""
    return {p.decode(): u.decode()
            for p, u in re.findall(rb'xmlns:([\w.-]+)\s*=\s*"([^"]*)"', xml)}


def _named_prefixes(xml: bytes) -> set[str]:
    """被 `Ignorable` / `Requires` 用名字點到的前綴。"""
    out: set[str] = set()
    for v in re.findall(rb'(?:Ignorable|Requires)\s*=\s*"([^"]*)"', xml):
        out |= set(v.decode().split())
    return out


def _check_part(original: bytes, written: bytes) -> list[str]:
    probs: list[str] = []
    try:
        fromstring(written)
    except Exception as e:  # noqa: BLE001
        return [f"讀不進去：{e}"]
    declared = _decls(written)
    missing = _named_prefixes(written) - declared.keys()
    if missing:
        probs.append(f"被點名但沒有宣告的前綴：{sorted(missing)}")
    if re.search(rb"</?ns\d+:|\sns\d+:", written):
        probs.append("前綴被改名成 nsN")
    uri_prefix_before = {u: p for p, u in _decls(original).items()}
    uri_prefix_after = {u: p for p, u in declared.items()}
    for uri, p in uri_prefix_before.items():
        if uri in uri_prefix_after and uri_prefix_after[uri] != p:
            probs.append(f"{uri} 的前綴從 {p} 變成 {uri_prefix_after[uri]}")
    return probs


# ---------- 試算表：客戶回報的那一種 ----------

def test_excel_saved_sheet_keeps_its_prefixes_and_declarations():
    data = _excel_like_xlsx()
    out, state = _translate(data, ".xlsx")
    z = zipfile.ZipFile(io.BytesIO(out))
    part = "xl/worksheets/sheet1.xml"
    written = z.read(part)
    probs = _check_part(state["raw"][part], written)
    assert not probs, f"{part}：{probs}"
    # xr2 / xr3 沒被用到也要留著宣告 —— Ignorable 點名了它們
    for p in ("x14ac", "xr", "xr2", "xr3"):
        assert f'xmlns:{p}="'.encode() in M._root_start_tag(written), (
            f"根元素上少了 xmlns:{p}（mc:Ignorable 點名了它）")


def test_standalone_declaration_is_kept():
    out, _ = _translate(_excel_like_xlsx(), ".xlsx")
    written = zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml")
    assert written.startswith(
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'), written[:80]


def test_translated_workbook_reads_back_with_the_translations():
    from openpyxl import load_workbook
    out, _ = _translate(_excel_like_xlsx(), ".xlsx")
    ws = load_workbook(io.BytesIO(out)).active
    values = {c.value for row in ws.iter_rows() for c in row if c.value}
    assert "譯：Quarterly Report" in values, values
    assert "譯：Security requirements for vendors" in values, values


# ---------- Word：同一個病，換一個形狀 ----------

def test_word_saved_document_keeps_ignorable_and_requires_prefixes():
    data = _word_like_docx()
    out, state = _translate(data, ".docx")
    part = "word/document.xml"
    written = zipfile.ZipFile(io.BytesIO(out)).read(part)
    probs = _check_part(state["raw"][part], written)
    assert not probs, f"{part}：{probs}"
    assert b'Requires="wps"' in written
    assert "譯：Text in a box".encode() in written
    assert written.startswith(
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>')


def test_every_real_word_sample_keeps_named_prefixes_declared():
    """用 Word 存出來的真實樣本再跑一次（`temp_pdfs/` 沒有就跳過 —— 素材不上 git）。"""
    from pathlib import Path
    root = Path(__file__).resolve().parents[1] / "temp_pdfs"
    samples = [p for p in sorted(root.rglob("*.docx")) + sorted(root.rglob("*.xlsx"))
               if not p.name.startswith("~$")]
    if not samples:
        pytest.skip("沒有 temp_pdfs 樣本（公開樹 / CI）")
    bad = []
    for p in samples:
        data = p.read_bytes()
        try:
            out, state = _translate(data, p.suffix.lower())
        except AssertionError:
            continue
        z = zipfile.ZipFile(io.BytesIO(out))
        for name in state["trees"]:
            probs = _check_part(state["raw"][name], z.read(name))
            if probs:
                bad.append((p.name, name, probs))
    assert not bad, bad[:5]


# ---------- 邊界 ----------

def test_one_prefix_bound_to_two_uris_does_not_produce_duplicate_declarations():
    """XML 允許內層把同一個前綴重新宣告成別的網址；提到根元素之後只能留一個。

    沒處理的話會輸出兩個 `xmlns:p=`，整份讀不進去（比前綴被改名更糟）。
    """
    doc = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_W}" xmlns:p="urn:example:one">'
        '<w:body><w:p p:a="1"><w:r><w:t>First paragraph here</w:t></w:r></w:p>'
        '<w:p xmlns:p="urn:example:two" p:b="2"><w:r><w:t>Second paragraph here</w:t>'
        '</w:r></w:p></w:body></w:document>'
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", doc)
    out, _ = _translate(buf.getvalue(), ".docx")
    written = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml")
    root = fromstring(written)            # 讀得進去 ＝ 沒有同名宣告
    ps = root.findall(f".//{{{_W}}}p")
    assert ps[0].get("{urn:example:one}a") == "1"
    assert ps[1].get("{urn:example:two}b") == "2", "第二個網址的屬性掉了或掛錯網址"
    assert written.count(b'xmlns:p="') == 1


def test_a_document_prefix_that_elementtree_already_uses_does_not_collide():
    """文件把 `dc` 拿去指別的網址，同時又用到 ElementTree 預設給 `dc` 的那個網址。

    後者如果只用預設命名空間宣告（沒有前綴），輸出時 ElementTree 會去查自己的預設表
    拿到 `dc` —— 跟文件的 `dc` 撞名，輸出兩個 `xmlns:dc=`，整份讀不進去。
    """
    dc = "http://purl.org/dc/elements/1.1/"
    doc = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_W}" xmlns:dc="urn:example:not-dublin-core">'
        '<w:body><w:p dc:note="1"><w:r><w:t>Collision check paragraph</w:t></w:r></w:p>'
        f'<meta xmlns="{dc}"/></w:body></w:document>'
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", doc)
    out, _ = _translate(buf.getvalue(), ".docx")
    written = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml")
    root = fromstring(written)                       # 讀得進去 ＝ 沒有同名宣告
    assert root.find(f".//{{{_W}}}p").get("{urn:example:not-dublin-core}note") == "1"
    assert root.find(f".//{{{dc}}}meta") is not None


def test_the_shared_namespace_map_is_restored():
    """`ET._namespace_map` 是全行程共用的 —— 用完不還原，會改到別的工具的輸出。

    素材用一個**只有這條測試會用**的網址，並且拿 `dc` 這個 ElementTree 預設就
    註冊給 Dublin Core 的前綴去指別的東西 —— 只拿前面測試用過的文件再跑一次的話，
    殘留的內容跟這次加進去的一模一樣，「前後相等」永遠成立（變異驗證抓到的）。
    """
    uri = "urn:example:jtdt-restore-check"
    dc = "http://purl.org/dc/elements/1.1/"
    assert ET._namespace_map.get(dc) == "dc", "前提：ElementTree 預設把 dc 給 Dublin Core"
    doc = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_W}" xmlns:dc="{uri}">'
        '<w:body><w:p dc:note="1"><w:r><w:t>Restore check paragraph</w:t></w:r></w:p>'
        '</w:body></w:document>'
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", doc)
    before = dict(ET._namespace_map)
    out, _ = _translate(buf.getvalue(), ".docx")
    written = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml")
    assert f'xmlns:dc="{uri}"'.encode() in written
    assert uri not in ET._namespace_map, "這份文件的前綴留在全行程共用的表裡了"
    assert ET._namespace_map.get(dc) == "dc", "Dublin Core 的預設前綴被弄丟了"
    assert dict(ET._namespace_map) == before


def test_concurrent_rebuilds_do_not_leak_prefixes_into_each_other():
    """兩個作業同時寫檔（作業佇列預設併行 2）：各自的前綴不可以互相污染。"""
    xlsx, docx = _excel_like_xlsx(), _word_like_docx()
    errors: list[str] = []

    def work(data, ext, part):
        for _ in range(15):
            out, state = _translate(data, ext)
            probs = _check_part(state["raw"][part],
                                zipfile.ZipFile(io.BytesIO(out)).read(part))
            if probs:
                errors.append(f"{ext}: {probs}")
                return

    ts = [threading.Thread(target=work, args=(xlsx, ".xlsx", "xl/worksheets/sheet1.xml")),
          threading.Thread(target=work, args=(docx, ".docx", "word/document.xml"))]
    for t in ts:
        t.start()
    for t in ts:
        t.join(timeout=60)
    assert not errors, errors[:3]
