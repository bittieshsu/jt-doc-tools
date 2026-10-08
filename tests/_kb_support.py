"""知識庫測試共用的東西：假的嵌入服務、把知識庫搬到暫存目錄、造測試文件。

**假的嵌入服務照兩種真的 API 的形狀寫**（Ollama `/api/embed`、OpenAI 相容
`/v1/embeddings`），向量是「相鄰兩字雜湊」算出來的 —— 字面越像向量越近，
所以檢索測試有意義，而且完全確定、不需要網路。

`kb_isolated`：只把**知識庫**的資料目錄換掉（`store.kb_dir` 與 `embed.settings_path`），
認證資料庫照用 conftest 那一份 —— 權限測試要用 conftest 的 `admin_session`
建使用者與群組。不去改 `settings.data_dir`（那會把作業佇列、認證全部一起搬走，
CLAUDE.md 記過合跑互相污染的事）。
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import threading
import unicodedata
import zlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest

DIM = 256


def fake_vec(text: str, dim: int = DIM, salt: str = "") -> list[float]:
    t = re.sub(r"\s+", "", unicodedata.normalize("NFKC", text or ""))
    v = [0.0] * dim
    for i in range(max(1, len(t) - 1)):
        g = (salt + t[i:i + 2]).encode("utf-8")
        v[zlib.crc32(g) % dim] += 1.0
    if not any(v):
        v[0] = 1.0
    return v


class FakeEmbed:
    """假的嵌入服務。

    * `kind`：`ollama`（只開 /api/embed，另有 GET /api/version）或 `openai`（只開 /v1/embeddings）
    * `fail_after`：第幾次請求之後一律回 500（模擬重建到一半服務掛掉）
    * `drift_after`：第幾次請求之後回不一樣的向量（模擬中途換了模型）
    * `context_limit`：單一輸入超過這麼多字就回 400「input length exceeds the context length」
    * `models`：模型清單（`ollama` 的 GET /api/tags 帶 `capabilities`；`openai` 的 GET /v1/models 只給名稱）
    """

    def __init__(self, kind: str = "ollama", *, dim: int = DIM, fail_after: int | None = None,
                 drift_after: int | None = None, context_limit: int | None = None,
                 status: int = 200, models: list[tuple[str, list[str]]] | None = None):
        self.kind = kind
        self.models = models if models is not None else [
            ("granite-embedding:278m", ["embedding"]), ("embeddinggemma:300m", ["embedding"]),
            ("gemma4:26b", ["completion", "tools"])]
        self.dim = dim
        self.fail_after = fail_after
        self.drift_after = drift_after
        self.context_limit = context_limit
        self.status = status
        self.requests: list[dict] = []
        self.auth: list[str] = []
        self.get_auth: list[str] = []
        self.paths: list[str] = []
        self._lock = threading.Lock()
        me = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):  # noqa: D401 — 安靜
                pass

            def do_GET(self):
                # 「是不是 Ollama」的探測（`LLMClient.is_ollama()` 打 /api/version）：Ollama 才有這支
                with me._lock:
                    me.paths.append("GET " + self.path)
                    me.get_auth.append(self.headers.get("Authorization", ""))
                if me.kind == "ollama" and self.path == "/api/version":
                    return self._send(200, {"version": "0.0-fake"})
                if me.kind == "ollama" and self.path == "/api/tags":
                    return self._send(200, {"models": [
                        {"name": n, "size": 1000, "capabilities": caps,
                         "details": {"embedding_length": me.dim, "context_length": 2048}}
                        for n, caps in me.models]})
                if me.kind == "openai" and self.path == "/v1/models":
                    return self._send(200, {"data": [{"id": n, "object": "model"} for n, _ in me.models]})
                return self._send(404, {"error": "not found"})

            def do_POST(self):
                n = int(self.headers.get("Content-Length") or 0)
                body = json.loads(self.rfile.read(n) or b"{}")
                with me._lock:
                    me.requests.append(body)
                    me.auth.append(self.headers.get("Authorization", ""))
                    me.paths.append(self.path)
                    count = len(me.requests)
                want = "/api/embed" if me.kind == "ollama" else "/v1/embeddings"
                if self.path != want:
                    return self._send(404, {"error": "not found"})
                if me.status != 200:
                    return self._send(me.status, {"error": "nope"})
                if me.fail_after is not None and count > me.fail_after:
                    return self._send(500, {"error": "backend crashed"})
                texts = body.get("input")
                if isinstance(texts, str):
                    texts = [texts]
                if me.context_limit and any(len(t) > me.context_limit for t in texts):
                    return self._send(400, {"error": "the input length exceeds the context length"})
                salt = "drift" if (me.drift_after is not None and count > me.drift_after) else ""
                vecs = [fake_vec(t, me.dim, salt) for t in texts]
                if me.kind == "ollama":
                    return self._send(200, {"model": body.get("model"), "embeddings": vecs})
                return self._send(200, {"data": [{"index": i, "embedding": v}
                                                  for i, v in enumerate(vecs)]})

            def _send(self, code, obj):
                raw = json.dumps(obj).encode("utf-8")
                self.send_response(code)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), H)
        self.port = self.server.server_address[1]

    @property
    def base(self) -> str:
        return f"http://127.0.0.1:{self.port}"

    def __enter__(self):
        self._t = threading.Thread(target=self.server.serve_forever, daemon=True)
        self._t.start()
        return self

    def __exit__(self, *exc):
        self.server.shutdown()
        self.server.server_close()


@pytest.fixture
def kb_isolated(tmp_path, monkeypatch):
    """知識庫用自己的暫存目錄（每條測試一份）。"""
    from app.core.kb import embed, store
    root = tmp_path / "kbdata"
    monkeypatch.setattr(store, "kb_dir", lambda: root / "knowledge")
    monkeypatch.setattr(embed, "settings_path", lambda: root / "knowledge_settings.json")
    yield root


# ---------------------------------------------------------------- 造文件
def make_pdf(pages: list[list[str]]) -> bytes:
    """每一頁一串行，寫進 PDF（抽得回來的那種）。

    用 PyMuPDF 內建的中文字型 `china-t`：畫出來的字形不重要（只驗抽字），
    而且不必找系統字型 —— 沒裝中日韓字型的 CI 也跑得動。
    """
    import fitz
    doc = fitz.open()
    for lines in pages:
        page = doc.new_page(width=595, height=842)
        y = 60
        for ln in lines:
            page.insert_text((60, y), ln, fontname="china-t", fontsize=11)
            y += 18
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_docx(paras: list[tuple[str, str | None]]) -> bytes:
    """[(文字, 樣式名或 None)] → .docx。"""
    import docx
    d = docx.Document()
    for text, style in paras:
        if style:
            d.add_paragraph(text, style=style)
        else:
            d.add_paragraph(text)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "欄一"
    t.rows[0].cells[1].text = "欄二"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_odt(paras: list[tuple[str, int | None]]) -> bytes:
    """[(文字, 標題層級或 None)] → 最小的 .odt。"""
    import zipfile
    body = []
    for text, lvl in paras:
        esc = text.replace("&", "&amp;").replace("<", "&lt;")
        if lvl:
            body.append(f'<text:h text:outline-level="{lvl}">{esc}</text:h>')
        else:
            body.append(f"<text:p>{esc}</text:p>")
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" office:version="1.3">'
        "<office:body><office:text>" + "".join(body) + "</office:text></office:body>"
        "</office:document-content>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text",
                    compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content)
    return buf.getvalue()


def add_doc(dataset_id: str, data: bytes, ext: str, title: str = "測試文件",
            *, activate: bool = True, process: bool = True) -> dict:
    """直接走 store ＋ indexer（不經背景作業），回版本。"""
    from app.core.kb import indexer, store
    v = store.create_version(dataset_id, title=title, filename="x" + ext, ext=ext, data=data,
                             sha256=hashlib.sha256(data).hexdigest(), meta={})
    if process:
        res = indexer.process_version(v["id"])
        assert res.get("ok"), res
        if activate:
            store.transition(v["id"], allowed_from=("ready",), to="active", activated_by="t")
    return store.get_version(v["id"])


HANDBOOK_TXT = """壹、總述
一、本手冊所稱文書，指處理公務或與公務有關之一切資料。
二、文書製作應採由左至右之橫行格式。
貳、公文製作
十八、公文用語規定如下：
（一）期望、目的及准駁用語，得視需要酌用「請」、「希」、「查照」、「鑒核」或「核示」等。
（二）直接稱謂用語：有隸屬關係之機關，上級對下級稱「貴」；下級對上級稱「鈞」；自稱「本」。
十九、簽之撰擬：簽應按「主旨」、「說明」、「擬辦」3段式辦理。
參、文書保密
六十二、機密文書對外發文時，應封裝於雙封套內，內封套左上角加蓋機密等級。
"""


def setup_embed(base: str, *, kind: str = "ollama", model: str = "fake-embed") -> None:
    """另外指定一台嵌入服務（測試用的假服務）；不沿用 LLM 伺服器。"""
    from app.core.kb import embed
    embed.save({"kind": kind, "base_url": base, "model": model, "use_llm_server": False,
                "batch_size": 4})


def read_json(p: Path) -> dict:
    return json.loads(p.read_text(encoding="utf-8"))
